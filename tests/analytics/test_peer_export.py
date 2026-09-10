from openpyxl import load_workbook

from src.analytics.peer_export import (
    OUTPUT_PATH,
    validate_workbook,
    write_workbook,
)


def test_peer_workbook_created():
    path = write_workbook()

    assert path.exists()


def test_exactly_11_sheets():
    write_workbook()

    wb = load_workbook(
        OUTPUT_PATH,
        data_only=True,
    )

    assert len(wb.sheetnames) == 11


def test_20_metric_columns():
    write_workbook()

    wb = load_workbook(
        OUTPUT_PATH,
        data_only=True,
    )

    for ws in wb.worksheets:

        # 3 identifier columns:
        # company_id
        # company_name
        # is_benchmark
        #
        # 20 metric columns:
        # 10 raw + 10 percentile
        assert ws.max_column == 23


def test_each_sheet_has_median():
    write_workbook()

    wb = load_workbook(
        OUTPUT_PATH,
        data_only=True,
    )

    for ws in wb.worksheets:
        assert (
            ws.cell(
                row=ws.max_row,
                column=1,
            ).value
            == "MEDIAN"
        )


def test_each_sheet_has_benchmark():
    write_workbook()

    wb = load_workbook(
        OUTPUT_PATH,
        data_only=True,
    )

    for ws in wb.worksheets:

        values = [
            ws.cell(
                row=row,
                column=3,
            ).value
            for row in range(
                2,
                ws.max_row,
            )
        ]

        assert any(
            value
            in (
                1,
                True,
                "1",
                "True",
            )
            for value in values
        )


def test_workbook_validation():
    write_workbook()

    assert validate_workbook() is True
