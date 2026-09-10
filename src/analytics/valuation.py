import sqlite3
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ============================================================
# PATHS
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parents[2]

DB_PATH = PROJECT_ROOT / "db" / "nifty100.db"

OUTPUT_DIR = PROJECT_ROOT / "output"

VALUATION_XLSX = OUTPUT_DIR / "valuation_summary.xlsx"

FLAGS_CSV = OUTPUT_DIR / "valuation_flags.csv"


# ============================================================
# LOAD SOURCE DATA
# ============================================================


def load_data():
    """Load data."""
    with sqlite3.connect(DB_PATH) as conn:

        companies = pd.read_sql_query(
            """
            SELECT
                id AS company_id,
                company_name
            FROM companies
            ORDER BY id
            """,
            conn,
        )

        sectors = pd.read_sql_query(
            """
            SELECT
                company_id,
                broad_sector AS sector
            FROM sectors
            """,
            conn,
        )

        market = pd.read_sql_query(
            """
            SELECT
                company_id,
                year,
                market_cap_crore,
                enterprise_value_crore,
                pe_ratio,
                pb_ratio,
                ev_ebitda,
                dividend_yield_pct
            FROM market_cap
            ORDER BY
                company_id,
                year
            """,
            conn,
        )

        ratios = pd.read_sql_query(
            """
            SELECT
                company_id,
                year,
                free_cash_flow_cr
            FROM financial_ratios
            ORDER BY
                company_id,
                year
            """,
            conn,
        )

    return (
        companies,
        sectors,
        market,
        ratios,
    )


# ============================================================
# LATEST MARKET YEAR
# ============================================================


def get_latest_market_year(
    market,
):
    """Return latest market year."""
    years = market["year"].dropna().astype(str).sort_values()

    if years.empty:
        raise ValueError("No market-cap years found.")

    return years.iloc[-1]


# ============================================================
# 5-YEAR MEDIAN P/E
# ============================================================


def calculate_5yr_median_pe(
    market,
    latest_year,
):
    """Calculate 5yr median pe."""
    market = market.copy()

    market["year_date"] = pd.to_datetime(
        market["year"],
        format="%Y-%m",
        errors="coerce",
    )

    latest_date = pd.to_datetime(
        latest_year,
        format="%Y-%m",
    )

    start_date = latest_date - pd.DateOffset(years=4)

    five_year = market[
        (market["year_date"] >= start_date) & (market["year_date"] <= latest_date)
    ].copy()

    result = (
        five_year.groupby(
            "company_id",
            as_index=False,
        )["pe_ratio"]
        .median()
        .rename(columns={"pe_ratio": "5yr_median_PE"})
    )

    return result


# ============================================================
# LATEST FCF
# Same financial period as latest market-cap year
# ============================================================


def get_latest_fcf(
    ratios,
    latest_year,
):
    """Return latest fcf."""
    latest = ratios[ratios["year"] == latest_year].copy()

    latest = latest.sort_values(
        [
            "company_id",
            "year",
        ]
    ).drop_duplicates(
        subset=["company_id"],
        keep="last",
    )

    return latest[
        [
            "company_id",
            "free_cash_flow_cr",
        ]
    ]


# ============================================================
# SECTOR MEDIAN P/E
# ============================================================


def calculate_sector_median_pe(
    latest_market,
    sectors,
):
    """Calculate sector median pe."""
    data = latest_market.merge(
        sectors,
        on="company_id",
        how="left",
    )

    medians = (
        data.groupby(
            "sector",
            dropna=False,
        )["pe_ratio"]
        .median()
        .reset_index(name="sector_median_PE")
    )

    return medians


# ============================================================
# VALUATION FLAG
# ============================================================


def valuation_flag(
    pe,
    sector_median,
):
    """Handle valuation flag."""
    if pd.isna(pe) or pd.isna(sector_median) or sector_median <= 0:
        return "N/A"

    if pe > (sector_median * 1.5):
        return "Caution"

    if pe < (sector_median * 0.7):
        return "Discount"

    return "Fair"


# ============================================================
# BUILD VALUATION DATAFRAME
# ============================================================


def build_valuation():
    """Build valuation."""
    (
        companies,
        sectors,
        market,
        ratios,
    ) = load_data()

    latest_year = get_latest_market_year(market)

    print(
        "Latest market year:",
        latest_year,
    )

    latest_market = market[market["year"] == latest_year].copy()

    latest_market = latest_market.sort_values("company_id").drop_duplicates(
        subset=["company_id"],
        keep="last",
    )

    five_year_pe = calculate_5yr_median_pe(
        market,
        latest_year,
    )

    latest_fcf = get_latest_fcf(
        ratios,
        latest_year,
    )

    sector_medians = calculate_sector_median_pe(
        latest_market,
        sectors,
    )

    df = companies.merge(
        sectors,
        on="company_id",
        how="left",
    )

    df = df.merge(
        latest_market,
        on="company_id",
        how="left",
    )

    df = df.merge(
        latest_fcf,
        on="company_id",
        how="left",
    )

    df = df.merge(
        five_year_pe,
        on="company_id",
        how="left",
    )

    df = df.merge(
        sector_medians,
        on="sector",
        how="left",
    )

    # --------------------------------------------------------
    # FCF YIELD
    # --------------------------------------------------------

    df["FCF_yield_pct"] = np.where(
        (df["market_cap_crore"].notna())
        & (df["market_cap_crore"] > 0)
        & (df["free_cash_flow_cr"].notna()),
        (df["free_cash_flow_cr"] / df["market_cap_crore"] * 100),
        np.nan,
    )

    # --------------------------------------------------------
    # P/E VS SECTOR MEDIAN %
    # --------------------------------------------------------

    df["PE_vs_sector_median_pct"] = np.where(
        (df["sector_median_PE"].notna()) & (df["sector_median_PE"] != 0),
        ((df["pe_ratio"] - df["sector_median_PE"]) / df["sector_median_PE"] * 100),
        np.nan,
    )

    # --------------------------------------------------------
    # FLAGS
    # --------------------------------------------------------

    df["flag"] = df.apply(
        lambda row: valuation_flag(
            row["pe_ratio"],
            row["sector_median_PE"],
        ),
        axis=1,
    )

    # --------------------------------------------------------
    # ROUNDING
    # --------------------------------------------------------

    numeric_columns = [
        "pe_ratio",
        "pb_ratio",
        "ev_ebitda",
        "FCF_yield_pct",
        "5yr_median_PE",
        "PE_vs_sector_median_pct",
        "sector_median_PE",
        "market_cap_crore",
        "free_cash_flow_cr",
    ]

    for column in numeric_columns:
        if column in df.columns:
            df[column] = pd.to_numeric(
                df[column],
                errors="coerce",
            ).round(2)

    return (
        df,
        latest_year,
    )


# ============================================================
# EXPORT EXCEL
# ============================================================


def export_excel(
    df,
):
    """Export excel."""
    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    export = df[
        [
            "company_id",
            "company_name",
            "sector",
            "pe_ratio",
            "pb_ratio",
            "ev_ebitda",
            "FCF_yield_pct",
            "5yr_median_PE",
            "PE_vs_sector_median_pct",
            "flag",
        ]
    ].copy()

    export = export.rename(
        columns={
            "pe_ratio": "P/E",
            "pb_ratio": "P/B",
            "ev_ebitda": "EV/EBITDA",
        }
    )

    export.to_excel(
        VALUATION_XLSX,
        index=False,
        sheet_name="Valuation Summary",
    )

    style_excel(VALUATION_XLSX)

    return export


# ============================================================
# EXPORT FLAG CSV
# ============================================================


def export_flags(
    df,
    latest_year,
):
    """Export flags."""
    flags = df[
        df["flag"].isin(
            [
                "Caution",
                "Discount",
            ]
        )
    ].copy()

    flags["latest_year"] = latest_year

    flags = flags[
        [
            "company_id",
            "company_name",
            "sector",
            "latest_year",
            "pe_ratio",
            "sector_median_PE",
            "5yr_median_PE",
            "PE_vs_sector_median_pct",
            "pb_ratio",
            "ev_ebitda",
            "market_cap_crore",
            "free_cash_flow_cr",
            "FCF_yield_pct",
            "flag",
        ]
    ]

    flags = flags.rename(
        columns={
            "pe_ratio": "P/E",
            "pb_ratio": "P/B",
            "ev_ebitda": "EV/EBITDA",
        }
    )

    flags.to_csv(
        FLAGS_CSV,
        index=False,
    )

    return flags


# ============================================================
# EXCEL FORMATTING
# ============================================================


def style_excel(
    path,
):
    """Style excel."""
    workbook = load_workbook(path)

    sheet = workbook["Valuation Summary"]

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    caution_fill = PatternFill(
        fill_type="solid",
        fgColor="F4CCCC",
    )

    discount_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAD3",
    )

    fair_fill = PatternFill(
        fill_type="solid",
        fgColor="FFF2CC",
    )

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    flag_column = None

    for cell in sheet[1]:
        if cell.value == "flag":
            flag_column = cell.column
            break

    if flag_column is not None:
        for row in range(
            2,
            sheet.max_row + 1,
        ):
            cell = sheet.cell(
                row=row,
                column=flag_column,
            )

            if cell.value == "Caution":
                cell.fill = caution_fill

            elif cell.value == "Discount":
                cell.fill = discount_fill

            elif cell.value == "Fair":
                cell.fill = fair_fill

    for column_cells in sheet.columns:
        max_length = 0

        column_letter = get_column_letter(column_cells[0].column)

        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)

            max_length = max(
                max_length,
                len(value),
            )

        sheet.column_dimensions[column_letter].width = min(
            max(
                max_length + 2,
                12,
            ),
            32,
        )

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    workbook.save(path)


# ============================================================
# VALIDATION
# ============================================================


def validate(
    full_df,
    summary,
    flags,
):
    """Handle validate."""
    required_columns = [
        "company_id",
        "company_name",
        "sector",
        "P/E",
        "P/B",
        "EV/EBITDA",
        "FCF_yield_pct",
        "5yr_median_PE",
        "PE_vs_sector_median_pct",
        "flag",
    ]

    assert len(summary) == 92, f"Expected 92 rows, " f"found {len(summary)}"

    assert summary["company_id"].nunique() == 92

    assert all(column in summary.columns for column in required_columns)

    assert summary["P/E"].notna().sum() == 92

    assert summary["5yr_median_PE"].notna().sum() == 92

    assert (
        summary["flag"]
        .isin(
            [
                "Caution",
                "Discount",
                "Fair",
                "N/A",
            ]
        )
        .all()
    )

    assert (
        flags["flag"]
        .isin(
            [
                "Caution",
                "Discount",
            ]
        )
        .all()
    )

    # Verify threshold logic
    check = full_df[full_df["sector_median_PE"].notna()].copy()

    caution = check[check["flag"] == "Caution"]

    if not caution.empty:
        assert (caution["pe_ratio"] > caution["sector_median_PE"] * 1.5).all()

    discount = check[check["flag"] == "Discount"]

    if not discount.empty:
        assert (discount["pe_ratio"] < discount["sector_median_PE"] * 0.7).all()


# ============================================================
# MAIN
# ============================================================


def main():
    """Run the module entry point."""
    full_df, latest_year = build_valuation()

    summary = export_excel(full_df)

    flags = export_flags(
        full_df,
        latest_year,
    )

    validate(
        full_df,
        summary,
        flags,
    )

    print()
    print("=" * 70)
    print("VALUATION MODULE COMPLETE")
    print("=" * 70)

    print(
        "Latest valuation year:",
        latest_year,
    )

    print(
        "Companies:",
        len(summary),
    )

    print(
        "FCF yield available:",
        summary["FCF_yield_pct"].notna().sum(),
    )

    print(
        "5Y median P/E available:",
        summary["5yr_median_PE"].notna().sum(),
    )

    print()
    print("Valuation flags:")

    print(summary["flag"].value_counts().to_string())

    print()
    print(
        "Flagged companies:",
        len(flags),
    )

    print()
    print(
        "Excel:",
        VALUATION_XLSX,
    )

    print(
        "CSV:",
        FLAGS_CSV,
    )

    print()
    print("DAY 26 VALUATION: PASS")


if __name__ == "__main__":
    main()
