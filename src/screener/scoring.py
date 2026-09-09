from pathlib import Path
import sqlite3

import numpy as np
import pandas as pd

from src.screener.engine import (
    DB_PATH,
    PROJECT_ROOT,
    load_screener_dataframe,
    run_all_presets,
)


OUTPUT_DIR = PROJECT_ROOT / "output"
OUTPUT_PATH = OUTPUT_DIR / "screener_output.xlsx"


# ============================================================
# SCORE WEIGHTS
# ============================================================

WEIGHTS = {
    "roe_score": 15,
    "roce_score": 10,
    "npm_score": 10,

    "fcf_cagr_score": 15,
    "cfo_pat_score": 10,
    "fcf_positive_score": 5,

    "revenue_cagr_score": 10,
    "pat_cagr_score": 10,

    "de_score": 10,
    "icr_score": 5,
}


# ============================================================
# FCF CAGR
# ============================================================

def calculate_fcf_cagr(
    start_value,
    end_value,
    years,
):
    if (
        pd.isna(start_value)
        or pd.isna(end_value)
        or years <= 0
    ):
        return np.nan

    # CAGR is not economically meaningful when
    # the starting FCF is zero or negative.
    if start_value <= 0:
        return np.nan

    # Positive-to-negative transition.
    if end_value <= 0:
        return np.nan

    return (
        (
            end_value / start_value
        ) ** (1 / years)
        - 1
    ) * 100


def load_fcf_history(
    db_path=DB_PATH,
):
    with sqlite3.connect(db_path) as conn:
        df = pd.read_sql_query(
            """
            SELECT
                company_id,
                year,
                free_cash_flow_cr
            FROM financial_ratios
            WHERE free_cash_flow_cr IS NOT NULL
            ORDER BY company_id, year
            """,
            conn,
        )

    df["period"] = pd.to_datetime(
        df["year"],
        format="%Y-%m",
        errors="coerce",
    )

    return df


def calculate_fcf_cagr_5yr(
    db_path=DB_PATH,
):
    history = load_fcf_history(
        db_path
    )

    rows = []

    for company_id, group in history.groupby(
        "company_id"
    ):
        group = (
            group[
                group["period"].notna()
            ]
            .sort_values("period")
            .reset_index(drop=True)
        )

        if group.empty:
            continue

        # Prefer latest March/full-year observation.
        march = group[
            group["period"].dt.month.eq(3)
        ]

        if not march.empty:
            latest = march.iloc[-1]
        else:
            latest = group.iloc[-1]

        target_period = (
            latest["period"]
            - pd.DateOffset(years=5)
        )

        previous = group[
            group["period"].eq(
                target_period
            )
        ]

        if previous.empty:
            cagr = np.nan
            start_year = None
            start_fcf = np.nan
        else:
            previous = previous.iloc[-1]

            start_fcf = previous[
                "free_cash_flow_cr"
            ]

            start_year = previous["year"]

            cagr = calculate_fcf_cagr(
                start_fcf,
                latest["free_cash_flow_cr"],
                5,
            )

        rows.append(
            {
                "company_id": company_id,
                "fcf_cagr_5yr": cagr,
                "fcf_cagr_start_year": start_year,
                "fcf_cagr_end_year": latest["year"],
                "fcf_cagr_start_value": start_fcf,
                "fcf_cagr_end_value": latest[
                    "free_cash_flow_cr"
                ],
            }
        )

    return pd.DataFrame(rows)


# ============================================================
# P10 / P90 WINSORIZATION
# ============================================================

def winsorized_score(
    series,
    inverse=False,
):
    numeric = pd.to_numeric(
        series,
        errors="coerce",
    )

    valid = numeric.dropna()

    result = pd.Series(
        np.nan,
        index=series.index,
        dtype=float,
    )

    if valid.empty:
        return result

    p10 = valid.quantile(0.10)
    p90 = valid.quantile(0.90)

    if np.isclose(
        p10,
        p90,
        equal_nan=False,
    ):
        result.loc[
            numeric.notna()
        ] = 50.0

        return result

    capped = numeric.clip(
        lower=p10,
        upper=p90,
    )

    scaled = (
        (capped - p10)
        / (p90 - p10)
        * 100
    )

    if inverse:
        scaled = 100 - scaled

    result.loc[
        numeric.notna()
    ] = scaled.loc[
        numeric.notna()
    ]

    return result.clip(
        lower=0,
        upper=100,
    )


# ============================================================
# COMPOSITE SCORE
# ============================================================

def calculate_composite_score(
    df,
):
    result = df.copy()

    result["roe_score"] = (
        winsorized_score(
            result[
                "return_on_equity_pct"
            ]
        )
    )

    result["roce_score"] = (
        winsorized_score(
            result[
                "return_on_capital_employed_pct"
            ]
        )
    )

    result["npm_score"] = (
        winsorized_score(
            result[
                "net_profit_margin_pct"
            ]
        )
    )

    result["fcf_cagr_score"] = (
        winsorized_score(
            result[
                "fcf_cagr_5yr"
            ]
        )
    )

    result["cfo_pat_score"] = (
        winsorized_score(
            result[
                "cfo_pat_ratio"
            ]
        )
    )

    result["fcf_positive_score"] = np.where(
        result[
            "free_cash_flow_cr"
        ].isna(),
        np.nan,
        np.where(
            result[
                "free_cash_flow_cr"
            ] > 0,
            100.0,
            0.0,
        ),
    )

    result["revenue_cagr_score"] = (
        winsorized_score(
            result[
                "revenue_cagr_5yr"
            ]
        )
    )

    result["pat_cagr_score"] = (
        winsorized_score(
            result[
                "pat_cagr_5yr"
            ]
        )
    )

    result["de_score"] = (
        winsorized_score(
            result[
                "debt_to_equity"
            ],
            inverse=True,
        )
    )

    icr = result[
        "interest_coverage_effective"
    ].replace(
        [np.inf, -np.inf],
        np.nan,
    )

    result["icr_score"] = (
        winsorized_score(icr)
    )

    # Debt-free companies get maximum ICR score.
    result.loc[
        np.isposinf(
            result[
                "interest_coverage_effective"
            ]
        ),
        "icr_score",
    ] = 100.0

    score_columns = list(
        WEIGHTS.keys()
    )

    weighted_sum = pd.Series(
        0.0,
        index=result.index,
    )

    available_weight = pd.Series(
        0.0,
        index=result.index,
    )

    for column, weight in WEIGHTS.items():
        valid = result[
            column
        ].notna()

        weighted_sum.loc[
            valid
        ] += (
            result.loc[
                valid,
                column,
            ]
            * weight
        )

        available_weight.loc[
            valid
        ] += weight

    result[
        "composite_quality_score"
    ] = np.where(
        available_weight > 0,
        weighted_sum
        / available_weight,
        np.nan,
    )

    result[
        "composite_quality_score"
    ] = (
        result[
            "composite_quality_score"
        ]
        .clip(0, 100)
        .round(2)
    )

    return result


# ============================================================
# SECTOR RELATIVE SCORE
# ============================================================

def add_sector_relative_score(
    df,
):
    result = df.copy()

    result[
        "sector_relative_composite_score"
    ] = np.nan

    for sector, group in result.groupby(
        "broad_sector",
        dropna=False,
    ):
        scores = group[
            "composite_quality_score"
        ]

        sector_score = (
            winsorized_score(scores)
        )

        result.loc[
            group.index,
            "sector_relative_composite_score",
        ] = sector_score

    result[
        "sector_relative_composite_score"
    ] = (
        result[
            "sector_relative_composite_score"
        ]
        .round(2)
    )

    return result


# ============================================================
# UPDATE DATABASE
# ============================================================

def update_database_scores(
    scored_df,
    db_path=DB_PATH,
):
    with sqlite3.connect(db_path) as conn:

        for _, row in scored_df.iterrows():

            conn.execute(
                """
                UPDATE financial_ratios
                SET composite_quality_score = ?
                WHERE company_id = ?
                  AND year = ?
                """,
                (
                    None
                    if pd.isna(
                        row[
                            "composite_quality_score"
                        ]
                    )
                    else float(
                        row[
                            "composite_quality_score"
                        ]
                    ),
                    row["company_id"],
                    row["year"],
                ),
            )

        conn.commit()


# ============================================================
# EXCEL EXPORT
# ============================================================

EXPORT_COLUMNS = [
    "company_id",
    "company_name",
    "broad_sector",
    "sub_sector",

    "return_on_equity_pct",
    "return_on_capital_employed_pct",
    "net_profit_margin_pct",

    "debt_to_equity",
    "interest_coverage",

    "free_cash_flow_cr",
    "fcf_cagr_5yr",
    "cfo_pat_ratio",

    "revenue_cagr_5yr",
    "pat_cagr_5yr",
    "eps_cagr_5yr",

    "operating_profit_margin_pct",
    "asset_turnover",

    "pe_ratio",
    "pb_ratio",
    "dividend_yield_pct",

    "market_cap_crore",
    "sales",
    "net_profit",

    "composite_quality_score",
    "sector_relative_composite_score",
]


def export_screener_workbook(
    scored_df,
    output_path=OUTPUT_PATH,
):
    from openpyxl import Workbook
    from openpyxl.styles import (
        Alignment,
        Font,
        PatternFill,
    )
    from openpyxl.utils import (
        get_column_letter,
    )

    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    # Re-run presets after DB composite
    # score update.
    preset_results = run_all_presets()

    score_lookup = scored_df[
        [
            "company_id",
            "fcf_cagr_5yr",
            "composite_quality_score",
            "sector_relative_composite_score",
        ]
    ].copy()

    wb = Workbook()

    default = wb.active
    wb.remove(default)

    green_fill = PatternFill(
        "solid",
        fgColor="C6EFCE",
    )

    red_fill = PatternFill(
        "solid",
        fgColor="FFC7CE",
    )

    yellow_fill = PatternFill(
        "solid",
        fgColor="FFEB9C",
    )

    header_fill = PatternFill(
        "solid",
        fgColor="D9EAF7",
    )

    header_font = Font(
        bold=True
    )

    preset_threshold_columns = {
        "quality_compounder": {
            "return_on_equity_pct": (
                "gt",
                15,
            ),
            "debt_to_equity": (
                "lt",
                1,
            ),
            "free_cash_flow_cr": (
                "gt",
                0,
            ),
            "revenue_cagr_5yr": (
                "gt",
                10,
            ),
        },

        "value_pick": {
            "pe_ratio": (
                "lt",
                20,
            ),
            "pb_ratio": (
                "lt",
                3,
            ),
            "debt_to_equity": (
                "lt",
                2,
            ),
            "dividend_yield_pct": (
                "gt",
                1,
            ),
        },

        "growth_accelerator": {
            "pat_cagr_5yr": (
                "gt",
                20,
            ),
            "revenue_cagr_5yr": (
                "gt",
                15,
            ),
            "debt_to_equity": (
                "lt",
                2,
            ),
        },

        "dividend_champion": {
            "dividend_yield_pct": (
                "gt",
                2,
            ),
            "dividend_payout_ratio_pct": (
                "lt",
                80,
            ),
            "free_cash_flow_cr": (
                "gt",
                0,
            ),
        },

        "debt_free_blue_chip": {
            "debt_to_equity": (
                "eq",
                0,
            ),
            "return_on_equity_pct": (
                "gt",
                12,
            ),
            "sales": (
                "gt",
                5000,
            ),
        },

        "turnaround_watch": {
            "revenue_cagr_3yr": (
                "gt",
                10,
            ),
            "free_cash_flow_cr": (
                "gt",
                0,
            ),
        },
    }

    for preset_name, preset_df in (
        preset_results.items()
    ):
        df = preset_df.copy()

        # Remove old score so merge does
        # not create duplicate columns.
        for col in [
            "composite_quality_score",
            "sector_relative_composite_score",
            "fcf_cagr_5yr",
        ]:
            if col in df.columns:
                df = df.drop(
                    columns=[col]
                )

        df = df.merge(
            score_lookup,
            on="company_id",
            how="left",
        )

        df = df.sort_values(
            "composite_quality_score",
            ascending=False,
            na_position="last",
        )

        extra_cols = []

        if (
            "dividend_payout_ratio_pct"
            in df.columns
        ):
            extra_cols.append(
                "dividend_payout_ratio_pct"
            )

        if (
            "revenue_cagr_3yr"
            in df.columns
        ):
            extra_cols.append(
                "revenue_cagr_3yr"
            )

        columns = []

        for col in (
            EXPORT_COLUMNS
            + extra_cols
        ):
            if (
                col in df.columns
                and col not in columns
            ):
                columns.append(col)

        export_df = df[
            columns
        ].copy()

        sheet_name = (
            preset_name
            .replace("_", " ")
            .title()
        )[:31]

        ws = wb.create_sheet(
            title=sheet_name
        )

        for col_idx, column in enumerate(
            export_df.columns,
            start=1,
        ):
            cell = ws.cell(
                row=1,
                column=col_idx,
                value=column,
            )

            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(
                horizontal="center"
            )

        for row_idx, row in enumerate(
            export_df.itertuples(
                index=False,
                name=None,
            ),
            start=2,
        ):
            for col_idx, value in enumerate(
                row,
                start=1,
            ):
                if pd.isna(value):
                    value = None

                ws.cell(
                    row=row_idx,
                    column=col_idx,
                    value=value,
                )

        thresholds = (
            preset_threshold_columns.get(
                preset_name,
                {},
            )
        )

        for column, (
            comparison,
            threshold,
        ) in thresholds.items():

            if column not in export_df.columns:
                continue

            col_idx = (
                list(
                    export_df.columns
                ).index(column)
                + 1
            )

            for row_idx in range(
                2,
                ws.max_row + 1,
            ):
                cell = ws.cell(
                    row=row_idx,
                    column=col_idx,
                )

                value = cell.value

                if value is None:
                    cell.fill = red_fill
                    continue

                # D/E is intentionally
                # skipped for Financials
                # in max-threshold presets.
                if (
                    column
                    == "debt_to_equity"
                    and preset_name
                    in {
                        "quality_compounder",
                        "value_pick",
                        "growth_accelerator",
                    }
                ):
                    sector_col = (
                        list(
                            export_df.columns
                        ).index(
                            "broad_sector"
                        )
                        + 1
                    )

                    sector = ws.cell(
                        row=row_idx,
                        column=sector_col,
                    ).value

                    if sector == "Financials":
                        cell.fill = yellow_fill
                        continue

                if comparison == "gt":
                    passed = (
                        value > threshold
                    )

                elif comparison == "lt":
                    passed = (
                        value < threshold
                    )

                elif comparison == "eq":
                    passed = np.isclose(
                        value,
                        threshold,
                        atol=1e-9,
                    )

                else:
                    passed = False

                cell.fill = (
                    green_fill
                    if passed
                    else red_fill
                )

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = (
            ws.dimensions
        )

        for idx, column in enumerate(
            export_df.columns,
            start=1,
        ):
            max_length = len(column)

            for row_idx in range(
                2,
                min(
                    ws.max_row,
                    100,
                ) + 1,
            ):
                value = ws.cell(
                    row=row_idx,
                    column=idx,
                ).value

                if value is not None:
                    max_length = max(
                        max_length,
                        len(str(value)),
                    )

            ws.column_dimensions[
                get_column_letter(idx)
            ].width = min(
                max_length + 2,
                24,
            )

    wb.save(output_path)

    return output_path


# ============================================================
# MAIN
# ============================================================

def main():
    print("=" * 80)
    print("SPRINT 3 — DAY 17")
    print("COMPOSITE QUALITY SCORE")
    print("=" * 80)

    df = load_screener_dataframe()

    fcf_cagr = calculate_fcf_cagr_5yr()

    df = df.merge(
        fcf_cagr,
        on="company_id",
        how="left",
    )

    df = calculate_composite_score(
        df
    )

    df = add_sector_relative_score(
        df
    )

    update_database_scores(df)

    output_path = (
        export_screener_workbook(df)
    )

    print()
    print("Companies scored:", len(df))

    print(
        "Composite non-null:",
        df[
            "composite_quality_score"
        ].notna().sum(),
    )

    print(
        "Sector score non-null:",
        df[
            "sector_relative_composite_score"
        ].notna().sum(),
    )

    print(
        "FCF CAGR 5Y non-null:",
        df[
            "fcf_cagr_5yr"
        ].notna().sum(),
    )

    print()
    print("TOP 10 QUALITY COMPANIES")
    print("-" * 80)

    top = (
        df[
            [
                "company_id",
                "broad_sector",
                "composite_quality_score",
                "sector_relative_composite_score",
                "return_on_equity_pct",
                "return_on_capital_employed_pct",
                "free_cash_flow_cr",
                "revenue_cagr_5yr",
                "pat_cagr_5yr",
            ]
        ]
        .sort_values(
            "composite_quality_score",
            ascending=False,
        )
        .head(10)
    )

    print(
        top.to_string(
            index=False
        )
    )

    print()
    print(
        "Workbook:",
        output_path,
    )

    print()
    print("DAY 17 COMPLETE")


if __name__ == "__main__":
    main()
